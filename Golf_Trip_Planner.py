import streamlit as st
import anthropic
import os
from tavily import TavilyClient
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import MergedCell
from dotenv import load_dotenv
import io

load_dotenv()

# Page config
st.set_page_config(
    page_title="Golf Trip Planner",
    page_icon="⛳",
    layout="centered"
)

st.sidebar.page_link("Golf_Trip_Planner.py", label="⛳ Trip Planner")
st.sidebar.page_link("pages/Gambling_Golf_Games.py", label="🏌️ Gambling Games")

# Custom CSS
st.markdown("""
<style>
    /* Main background */
    .stApp {
        background-color: #f8faf8;
    }
    
    /* Hero banner */
    .hero {
        background: linear-gradient(135deg, #1a5c2a 0%, #2d8a47 50%, #1a5c2a 100%);
        padding: 40px 30px;
        border-radius: 16px;
        text-align: center;
        margin-bottom: 30px;
        color: white;
    }
    .hero h1 {
        font-size: 2.4em;
        font-weight: 800;
        margin: 0;
        color: white;
    }
    .hero p {
        font-size: 1.1em;
        opacity: 0.9;
        margin-top: 10px;
        color: white;
    }

    /* Section headers */
    .section-header {
        background: linear-gradient(90deg, #1a5c2a, #2d8a47);
        color: white;
        padding: 12px 20px;
        border-radius: 10px;
        font-size: 1.1em;
        font-weight: 700;
        margin: 25px 0 15px 0;
    }

    /* Cards */
    .card {
        background: white;
        border-radius: 12px;
        padding: 20px;
        margin-bottom: 15px;
        border-left: 5px solid #2d8a47;
        box-shadow: 0 2px 8px rgba(0,0,0,0.06);
    }
    .card h3 {
        margin: 0 0 10px 0;
        color: #1a5c2a;
        font-size: 1.1em;
    }
    .card p {
        margin: 4px 0;
        color: #444;
        font-size: 0.95em;
    }

    /* Metric cards */
    .metric-card {
        background: white;
        border-radius: 10px;
        padding: 15px;
        text-align: center;
        box-shadow: 0 2px 8px rgba(0,0,0,0.06);
        margin-bottom: 10px;
    }
    .metric-label {
        font-size: 0.8em;
        color: #888;
        text-transform: uppercase;
        letter-spacing: 0.5px;
    }
    .metric-value {
        font-size: 1.8em;
        font-weight: 700;
        color: #1a5c2a;
    }

    /* Form styling */
    .stTextInput > div > div > input {
        border-radius: 8px;
        border: 1.5px solid #ddd;
        padding: 10px;
    }
    .stTextInput > div > div > input:focus {
        border-color: #2d8a47;
    }

    /* Button */
    .stButton > button {
        background: linear-gradient(135deg, #1a5c2a, #2d8a47);
        color: white;
        border: none;
        border-radius: 10px;
        padding: 14px 28px;
        font-size: 1.1em;
        font-weight: 700;
        width: 100%;
        cursor: pointer;
        margin-top: 10px;
    }
    .stButton > button:hover {
        background: linear-gradient(135deg, #154d23, #267a3e);
        transform: translateY(-1px);
    }

    /* Download button */
    .stDownloadButton > button {
        background: white;
        color: #1a5c2a;
        border: 2px solid #1a5c2a;
        border-radius: 10px;
        padding: 12px 24px;
        font-size: 1em;
        font-weight: 600;
        width: 100%;
    }

    /* Success box */
    .success-box {
        background: #e8f5e9;
        border: 1.5px solid #2d8a47;
        border-radius: 10px;
        padding: 15px 20px;
        text-align: center;
        color: #1a5c2a;
        font-weight: 600;
        margin: 20px 0;
    }

    /* Divider */
    .custom-divider {
        border: none;
        border-top: 2px solid #e8f5e9;
        margin: 20px 0;
    }

    /* GolfNow link */
    a {
        color: #2d8a47 !important;
        font-weight: 600;
    }

    /* Hide streamlit branding */
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    [data-testid="stSidebarNav"] {display: none !important;}
</style>
""", unsafe_allow_html=True)

# Initialize clients
tavily = TavilyClient(api_key=os.environ.get("TAVILY_API_KEY"))
client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"))

# Hero section
st.markdown("""
<div class="hero">
    <h1>⛳ Golf Trip Planner</h1>
    <p>Enter your trip details and we'll build a complete cost estimate — courses, flights, hotels, and more.</p>
</div>
""", unsafe_allow_html=True)

# Input section
st.markdown('<div class="section-header">📋 Trip Details</div>', unsafe_allow_html=True)

city = st.text_input("📍 Where are you planning to golf?", placeholder="e.g. Scottsdale, AZ")
dates = st.text_input("📅 What are the dates of your trip?", placeholder="e.g. Oct 11-13")
nights = st.number_input("🌙 How many nights?", min_value=1, max_value=7, value=2)
golfers = st.text_input("👥 How many golfers?", placeholder="e.g. 4")

st.markdown('<div class="section-header">✈️ Travel Preferences</div>', unsafe_allow_html=True)
transport = st.radio("How are you getting there?", ["Drive", "Fly"], horizontal=True)
departure = ""
if transport == "Fly":
    departure = st.text_input("🛫 What city are you flying from?", placeholder="e.g. Chicago, IL")

st.markdown('<div class="section-header">🏨 Accommodations</div>', unsafe_allow_html=True)
hotel = st.radio("Include hotel pricing?", ["Yes", "No"], horizontal=True)

st.markdown("---")
plan_button = st.button("⛳ Build My Golf Trip")

if plan_button:
    if not city or not dates or not golfers:
        st.error("⚠️ Please fill out all required fields before searching.")
    else:
        # Golf search
        st.markdown('<div class="section-header">⛳ Golf Courses</div>', unsafe_allow_html=True)
        with st.spinner("🔍 Searching for the best courses..."):
            search1 = tavily.search(f"golf courses weekend green fees pricing {city} 2026")
            search2 = tavily.search(f"best public golf courses {city} tee times rates")
            search_text = "\n".join([r["content"] for r in search1["results"] + search2["results"]])

            prompt = f"""
            Based on the following real web search results, put together a golf trip plan for {city} from {dates} for {golfers} golfers.

            SEARCH RESULTS:
            {search_text}

            Please provide exactly 5 golf courses with the following info for each:
            - Course name
            - Weekend green fee per person (low estimate)
            - Weekend green fee per person (high estimate)
            - One sentence description

            Format your response exactly like this for each course:
            COURSE: [name]
            LOW: [number only]
            HIGH: [number only]
            DESC: [one sentence]
            """

            message = client.messages.create(
                model="claude-sonnet-4-6",
                max_tokens=2048,
                messages=[{"role": "user", "content": prompt}]
            )
            response_text = message.content[0].text

        # Parse courses
        lines = response_text.split("\n")
        courses = []
        current = {}
        for line in lines:
            if line.startswith("COURSE:"):
                if current:
                    courses.append(current)
                current = {"name": line.replace("COURSE:", "").strip()}
            elif line.startswith("LOW:"):
                try:
                    current["low"] = int(line.replace("LOW:", "").strip())
                except:
                    current["low"] = 0
            elif line.startswith("HIGH:"):
                try:
                    current["high"] = int(line.replace("HIGH:", "").strip())
                except:
                    current["high"] = 0
            elif line.startswith("DESC:"):
                current["desc"] = line.replace("DESC:", "").strip()
        if current:
            courses.append(current)

        golfers_int = int(golfers)
        for course in courses:
            course_search = course['name'].replace(' ', '+')
            golfnow_url = f"https://www.golfnow.com/tee-times/search#search/facility-name={course_search}"
            st.markdown(f"""
            <div class="card">
                <h3>🏌️ {course['name']}</h3>
                <p>💰 <strong>Green Fee:</strong> ${course['low']}–${course['high']} per person</p>
                <p>👥 <strong>Group Total:</strong> ${course['low'] * golfers_int:,}–${course['high'] * golfers_int:,}</p>
                <p>📝 {course['desc']}</p>
                <p><a href="{golfnow_url}" target="_blank">⛳ Check Tee Times on GolfNow →</a></p>
            </div>
            """, unsafe_allow_html=True)

        # Flights
        flight_info = ""
        flight_data = {}
        if transport == "Fly" and departure:
            st.markdown('<div class="section-header">✈️ Flights & Rental Car</div>', unsafe_allow_html=True)
            with st.spinner("🔍 Searching for flights and rental cars..."):
                flight_search = tavily.search(f"flights from {departure} to {city} {dates} 2026 price")
                rental_search = tavily.search(f"rental car {city} airport {dates} 2026 average price")
                flight_text = "\n".join([r["content"] for r in flight_search["results"]])
                rental_text = "\n".join([r["content"] for r in rental_search["results"]])

                flight_prompt = f"""
                Based on these search results, provide estimates for:
                1. Round trip flight cost per person from {departure} to {city} for {dates}
                2. Rental car cost for the trip (total, not per person)

                FLIGHT SEARCH RESULTS:
                {flight_text}

                RENTAL CAR SEARCH RESULTS:
                {rental_text}

                Format exactly like this:
                FLIGHT_LOW: [number only]
                FLIGHT_HIGH: [number only]
                RENTAL_LOW: [number only]
                RENTAL_HIGH: [number only]
                FLIGHT_DESC: [one sentence summary]
                RENTAL_DESC: [one sentence summary]
                """

                flight_message = client.messages.create(
                    model="claude-sonnet-4-6",
                    max_tokens=1024,
                    messages=[{"role": "user", "content": flight_prompt}]
                )
                flight_info = flight_message.content[0].text

            for line in flight_info.split("\n"):
                for key in ["FLIGHT_LOW", "FLIGHT_HIGH", "RENTAL_LOW", "RENTAL_HIGH", "FLIGHT_DESC", "RENTAL_DESC"]:
                    if line.startswith(f"{key}:"):
                        flight_data[key] = line.replace(f"{key}:", "").strip()

            st.markdown(f"""
            <div class="card">
                <h3>✈️ Flights — {departure} to {city}</h3>
                <p>💰 <strong>Per Person:</strong> ${flight_data.get('FLIGHT_LOW', 'N/A')}–${flight_data.get('FLIGHT_HIGH', 'N/A')} round trip</p>
                <p>📝 {flight_data.get('FLIGHT_DESC', '').replace('$', '\\$')}</p>
            </div>
            <div class="card">
                <h3>🚗 Rental Car</h3>
                <p>💰 <strong>Total for Group:</strong> ${flight_data.get('RENTAL_LOW', 'N/A')}–${flight_data.get('RENTAL_HIGH', 'N/A')}</p>
                <p>📝 {flight_data.get('RENTAL_DESC', '').replace('$', '\\$')}</p>
            </div>
            """, unsafe_allow_html=True)

        # Hotels
        hotel_info = ""
        hotels = []
        if hotel == "Yes":
            st.markdown('<div class="section-header">🏨 Hotels</div>', unsafe_allow_html=True)
            with st.spinner("🔍 Searching for hotels..."):
                hotel_search = tavily.search(f"hotels near golf courses {city} {dates} price per night 2026")
                hotel_text = "\n".join([r["content"] for r in hotel_search["results"]])

                hotel_prompt = f"""
                Based on these search results, provide 3 hotel options near golf courses in {city} for {dates}.

                SEARCH RESULTS:
                {hotel_text}

                Format exactly like this for each:
                HOTEL: [name]
                HOTEL_LOW: [number only, per night]
                HOTEL_HIGH: [number only, per night]
                HOTEL_DESC: [one sentence]
                """

                hotel_message = client.messages.create(
                    model="claude-sonnet-4-6",
                    max_tokens=1024,
                    messages=[{"role": "user", "content": hotel_prompt}]
                )
                hotel_info = hotel_message.content[0].text

            current_hotel = {}
            for line in hotel_info.split("\n"):
                if line.startswith("HOTEL:"):
                    if current_hotel:
                        hotels.append(current_hotel)
                    current_hotel = {"name": line.replace("HOTEL:", "").strip()}
                elif line.startswith("HOTEL_LOW:"):
                    current_hotel["low"] = line.replace("HOTEL_LOW:", "").strip()
                elif line.startswith("HOTEL_HIGH:"):
                    current_hotel["high"] = line.replace("HOTEL_HIGH:", "").strip()
                elif line.startswith("HOTEL_DESC:"):
                    current_hotel["desc"] = line.replace("HOTEL_DESC:", "").strip()
            if current_hotel:
                hotels.append(current_hotel)

            for h in hotels:
                try:
                    nights_int = int(nights)
                    low_total = int(h.get('low', 0)) * nights_int
                    high_total = int(h.get('high', 0)) * nights_int
                    total_str = f"<p>👥 <strong>Total ({nights_int} nights):</strong> ${low_total:,}–${high_total:,}</p>"
                except:
                    total_str = ""

                st.markdown(f"""
                <div class="card">
                    <h3>🏨 {h['name']}</h3>
                    <p>💰 <strong>Per Night:</strong> ${h.get('low', 'N/A')}–${h.get('high', 'N/A')}</p>
                    {total_str}
                    <p>📝 {h.get('desc', '')}</p>
                </div>
                """, unsafe_allow_html=True)

        # Excel export
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Golf Trip"
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(fill_type="solid", fgColor="2E7D32")

        ws["A1"] = f"Golf Trip to {city} | {dates} | {golfers} Golfers"
        ws["A1"].font = Font(bold=True, size=16)
        ws.merge_cells("A1:F1")

        headers = ["Course", "Low Fee/Person", "High Fee/Person", "Low Total (Group)", "High Total (Group)", "Description"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row, course in enumerate(courses, 4):
            ws.cell(row=row, column=1, value=course.get("name", ""))
            ws.cell(row=row, column=2, value=f"${course.get('low', 0)}")
            ws.cell(row=row, column=3, value=f"${course.get('high', 0)}")
            ws.cell(row=row, column=4, value=f"${course.get('low', 0) * golfers_int}")
            ws.cell(row=row, column=5, value=f"${course.get('high', 0) * golfers_int}")
            ws.cell(row=row, column=6, value=course.get("desc", ""))

        current_row = len(courses) + 5

        if flight_info:
            ws.cell(row=current_row, column=1, value="FLIGHTS & RENTAL CAR").font = Font(bold=True, size=13)
            current_row += 1
            for col, header in enumerate(["", "Low Estimate", "High Estimate", "", "", "Notes"], 1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            current_row += 1
            ws.cell(row=current_row, column=1, value=f"Flights ({departure} to {city}) per person")
            ws.cell(row=current_row, column=2, value=f"${flight_data.get('FLIGHT_LOW', 'N/A')}")
            ws.cell(row=current_row, column=3, value=f"${flight_data.get('FLIGHT_HIGH', 'N/A')}")
            ws.cell(row=current_row, column=6, value=flight_data.get('FLIGHT_DESC', ''))
            current_row += 1
            ws.cell(row=current_row, column=1, value="Rental Car (total for group)")
            ws.cell(row=current_row, column=2, value=f"${flight_data.get('RENTAL_LOW', 'N/A')}")
            ws.cell(row=current_row, column=3, value=f"${flight_data.get('RENTAL_HIGH', 'N/A')}")
            ws.cell(row=current_row, column=6, value=flight_data.get('RENTAL_DESC', ''))
            current_row += 2

        if hotels:
            ws.cell(row=current_row, column=1, value="HOTELS").font = Font(bold=True, size=13)
            current_row += 1
            for col, header in enumerate(["Hotel", "Low/Night", "High/Night", "", "", "Description"], 1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            current_row += 1
            for h in hotels:
                ws.cell(row=current_row, column=1, value=h.get("name", ""))
                ws.cell(row=current_row, column=2, value=f"${h.get('low', 'N/A')}")
                ws.cell(row=current_row, column=3, value=f"${h.get('high', 'N/A')}")
                ws.cell(row=current_row, column=6, value=h.get("desc", ""))
                current_row += 1

        for col_idx in range(1, 7):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 60 if col_idx == 6 else (35 if col_idx == 1 else 20)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        st.markdown('<div class="success-box">✅ Your golf trip plan is ready!</div>', unsafe_allow_html=True)
        st.download_button(
            label="📥 Download Full Trip Report (Excel)",
            data=buffer,
            file_name=f"golf_trip_{city.replace(' ', '_')}_{dates.replace(' ', '_')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )