from dotenv import load_dotenv
load_dotenv()
from openpyxl.cell.cell import MergedCell
import anthropic
import os
from tavily import TavilyClient
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

tavily = TavilyClient(api_key=os.environ.get("TAVILY_API_KEY"))
client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"))

city = input("What city are you considering for your golf trip? ")
dates = input("What dates are you looking at? (e.g. Oct 11-13) ")
golfers = input("How many golfers? ")

print("\nSearching the web for current golf prices...\n")

search1 = tavily.search(f"golf courses weekend green fees pricing {city} 2026")
search2 = tavily.search(f"best public golf courses {city} tee times rates")
search_text = "\n".join([r["content"] for r in search1["results"] + search2["results"]])

prompt = f"""
Based on the following real web search results, put together a golf trip plan for {city} from {dates} for {golfers} golfers.

SEARCH RESULTS:
{search_text}

Please provide exactly 5 golf courses with the following info for each:
- Course name
- Weekend green fee per person (low estimate)
- Weekend green fee per person (high estimate)
- One sentence description

Then provide a total cost summary for the group.

Format your response exactly like this for each course:
COURSE: [name]
LOW: [number only]
HIGH: [number only]
DESC: [one sentence]
"""

message = client.messages.create(
    model="claude-sonnet-4-6",
    max_tokens=2048,
    messages=[{"role": "user", "content": prompt}]
)

response_text = message.content[0].text
print(response_text)

# Parse response and build Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Golf Trip"

# Header styling
header_font = Font(bold=True, color="FFFFFF", size=12)
header_fill = PatternFill(fill_type="solid", fgColor="2E7D32")
title_font = Font(bold=True, size=14)

# Title
ws["A1"] = f"Golf Trip to {city} | {dates} | {golfers} Golfers"
ws["A1"].font = Font(bold=True, size=16)
ws.merge_cells("A1:F1")

# Column headers
headers = ["Course", "Low Fee/Person", "High Fee/Person", "Low Total (Group)", "High Total (Group)", "Description"]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# Parse and write course data
lines = response_text.split("\n")
courses = []
current = {}

for line in lines:
    if line.startswith("COURSE:"):
        if current:
            courses.append(current)
        current = {"name": line.replace("COURSE:", "").strip()}
    elif line.startswith("LOW:"):
        try:
            current["low"] = int(line.replace("LOW:", "").strip())
        except:
            current["low"] = 0
    elif line.startswith("HIGH:"):
        try:
            current["high"] = int(line.replace("HIGH:", "").strip())
        except:
            current["high"] = 0
    elif line.startswith("DESC:"):
        current["desc"] = line.replace("DESC:", "").strip()

if current:
    courses.append(current)

for row, course in enumerate(courses, 4):
    golfers_int = int(golfers)
    ws.cell(row=row, column=1, value=course.get("name", ""))
    ws.cell(row=row, column=2, value=f"${course.get('low', 0)}")
    ws.cell(row=row, column=3, value=f"${course.get('high', 0)}")
    ws.cell(row=row, column=4, value=f"${course.get('low', 0) * golfers_int}")
    ws.cell(row=row, column=5, value=f"${course.get('high', 0) * golfers_int}")
    ws.cell(row=row, column=6, value=course.get("desc", ""))

# Auto size columns
for col in ws.iter_cols():
    max_length = max((len(str(cell.value)) for cell in col if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell)), default=10)
    if not isinstance(col[0], openpyxl.cell.cell.MergedCell):
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 50)

filename = f"golf_trip_{city.replace(' ', '_')}_{dates.replace(' ', '_')}.xlsx"
wb.save(filename)
print(f"\n✅ Spreadsheet saved as: {filename}")
input()