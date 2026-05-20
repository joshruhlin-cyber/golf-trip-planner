import anthropic
import os
from tavily import TavilyClient
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import MergedCell
from dotenv import load_dotenv

load_dotenv()

tavily = TavilyClient(api_key=os.environ.get("TAVILY_API_KEY"))
client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"))

city = input("What city are you considering for your golf trip? ")
dates = input("What dates are you looking at? (e.g. Oct 11-13) ")
golfers = input("How many golfers? ")

print("\nSearching the web for current golf prices...\n")

search1 = tavily.search(f"golf courses weekend green fees pricing {city} 2026")
search2 = tavily.search(f"best public golf courses {city} tee times rates")
search_text = "\n".join([r["content"] for r in search1["results"] + search2["results"]])

prompt = f"""
Based on the following real web search results, put together a golf trip plan for {city} from {dates} for {golfers} golfers.

SEARCH RESULTS:
{search_text}

Please provide exactly 5 golf courses with the following info for each:
- Course name
- Weekend green fee per person (low estimate)
- Weekend green fee per person (high estimate)
- One sentence description

Format your response exactly like this for each course:
COURSE: [name]
LOW: [number only]
HIGH: [number only]
DESC: [one sentence]
"""

message = client.messages.create(
    model="claude-sonnet-4-6",
    max_tokens=2048,
    messages=[{"role": "user", "content": prompt}]
)

response_text = message.content[0].text
print(response_text)
print("DEBUG: Golf search complete")

# Ask about transportation
print("\n" + "="*50)
print("✅ Golf pricing complete! Now let's talk about travel.")
print("="*50)
transport = input("\nWill you be driving or flying to the trip? Type 'drive' or 'fly': ").strip().lower()

flight_info = ""
rental_info = ""

if transport == "fly":
    departure = input("What city will you be flying from? ")
    
    print("\nSearching for flight and rental car prices...\n")
    
    flight_search = tavily.search(f"flights from {departure} to {city} {dates} 2026 price")
    rental_search = tavily.search(f"rental car {city} airport {dates} 2026 average price")
    
    flight_text = "\n".join([r["content"] for r in flight_search["results"]])
    rental_text = "\n".join([r["content"] for r in rental_search["results"]])
    
    flight_prompt = f"""
    Based on these search results, provide a brief estimate for:
    1. Round trip flight cost per person from {departure} to {city} for {dates}
    2. Rental car cost for the trip (total, not per person)
    
    FLIGHT SEARCH RESULTS:
    {flight_text}
    
    RENTAL CAR SEARCH RESULTS:
    {rental_text}
    
    Format exactly like this:
    FLIGHT_LOW: [number only]
    FLIGHT_HIGH: [number only]
    RENTAL_LOW: [number only]
    RENTAL_HIGH: [number only]
    FLIGHT_DESC: [one sentence summary]
    RENTAL_DESC: [one sentence summary]
    """
    
    flight_message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=1024,
        messages=[{"role": "user", "content": flight_prompt}]
    )
    
    flight_info = flight_message.content[0].text
    print(flight_info)

# Ask about hotels
print("\n" + "="*50)
hotel = input("\nWould you like hotel pricing too? (yes/no): ").strip().lower()

hotel_info = ""

if hotel == "yes":
    print("\nSearching for hotel prices...\n")
    
    hotel_search = tavily.search(f"hotels near golf courses {city} {dates} price per night 2026")
    hotel_text = "\n".join([r["content"] for r in hotel_search["results"]])
    
    hotel_prompt = f"""
    Based on these search results, provide 3 hotel options near golf courses in {city} for {dates}.
    
    SEARCH RESULTS:
    {hotel_text}
    
    Format exactly like this for each:
    HOTEL: [name]
    HOTEL_LOW: [number only, per night]
    HOTEL_HIGH: [number only, per night]
    HOTEL_DESC: [one sentence]
    """
    
    hotel_message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=1024,
        messages=[{"role": "user", "content": hotel_prompt}]
    )
    
    hotel_info = hotel_message.content[0].text
    print(hotel_info)

# Build Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Golf Trip"

header_font = Font(bold=True, color="FFFFFF", size=12)
header_fill = PatternFill(fill_type="solid", fgColor="2E7D32")

ws["A1"] = f"Golf Trip to {city} | {dates} | {golfers} Golfers"
ws["A1"].font = Font(bold=True, size=16)
ws.merge_cells("A1:F1")

headers = ["Course", "Low Fee/Person", "High Fee/Person", "Low Total (Group)", "High Total (Group)", "Description"]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

lines = response_text.split("\n")
courses = []
current = {}

for line in lines:
    if line.startswith("COURSE:"):
        if current:
            courses.append(current)
        current = {"name": line.replace("COURSE:", "").strip()}
    elif line.startswith("LOW:"):
        try:
            current["low"] = int(line.replace("LOW:", "").strip())
        except:
            current["low"] = 0
    elif line.startswith("HIGH:"):
        try:
            current["high"] = int(line.replace("HIGH:", "").strip())
        except:
            current["high"] = 0
    elif line.startswith("DESC:"):
        current["desc"] = line.replace("DESC:", "").strip()

if current:
    courses.append(current)

for row, course in enumerate(courses, 4):
    golfers_int = int(golfers)
    ws.cell(row=row, column=1, value=course.get("name", ""))
    ws.cell(row=row, column=2, value=f"${course.get('low', 0)}")
    ws.cell(row=row, column=3, value=f"${course.get('high', 0)}")
    ws.cell(row=row, column=4, value=f"${course.get('low', 0) * golfers_int}")
    ws.cell(row=row, column=5, value=f"${course.get('high', 0) * golfers_int}")
    ws.cell(row=row, column=6, value=course.get("desc", ""))

current_row = len(courses) + 5

# Add flight and rental info if flying
if flight_info:
    ws.cell(row=current_row, column=1, value="FLIGHTS & RENTAL CAR")
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=13)
    current_row += 1
    
    flight_headers = ["", "Low Estimate", "High Estimate", "", "", "Notes"]
    for col, header in enumerate(flight_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    current_row += 1

    flight_lines = flight_info.split("\n")
    flight_data = {}
    for line in flight_lines:
        for key in ["FLIGHT_LOW", "FLIGHT_HIGH", "RENTAL_LOW", "RENTAL_HIGH", "FLIGHT_DESC", "RENTAL_DESC"]:
            if line.startswith(f"{key}:"):
                flight_data[key] = line.replace(f"{key}:", "").strip()

    ws.cell(row=current_row, column=1, value=f"Flights ({departure} to {city}) per person")
    ws.cell(row=current_row, column=2, value=f"${flight_data.get('FLIGHT_LOW', 'N/A')}")
    ws.cell(row=current_row, column=3, value=f"${flight_data.get('FLIGHT_HIGH', 'N/A')}")
    ws.cell(row=current_row, column=6, value=flight_data.get('FLIGHT_DESC', ''))
    current_row += 1

    ws.cell(row=current_row, column=1, value="Rental Car (total for group)")
    ws.cell(row=current_row, column=2, value=f"${flight_data.get('RENTAL_LOW', 'N/A')}")
    ws.cell(row=current_row, column=3, value=f"${flight_data.get('RENTAL_HIGH', 'N/A')}")
    ws.cell(row=current_row, column=6, value=flight_data.get('RENTAL_DESC', ''))
    current_row += 2

# Add hotel info
if hotel_info:
    ws.cell(row=current_row, column=1, value="HOTELS")
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=13)
    current_row += 1

    hotel_headers = ["Hotel", "Low/Night", "High/Night", "", "", "Description"]
    for col, header in enumerate(hotel_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    current_row += 1

    hotel_lines = hotel_info.split("\n")
    hotels = []
    current_hotel = {}

    for line in hotel_lines:
        if line.startswith("HOTEL:"):
            if current_hotel:
                hotels.append(current_hotel)
            current_hotel = {"name": line.replace("HOTEL:", "").strip()}
        elif line.startswith("HOTEL_LOW:"):
            current_hotel["low"] = line.replace("HOTEL_LOW:", "").strip()
        elif line.startswith("HOTEL_HIGH:"):
            current_hotel["high"] = line.replace("HOTEL_HIGH:", "").strip()
        elif line.startswith("HOTEL_DESC:"):
            current_hotel["desc"] = line.replace("HOTEL_DESC:", "").strip()

    if current_hotel:
        hotels.append(current_hotel)

    for hotel_row in hotels:
        ws.cell(row=current_row, column=1, value=hotel_row.get("name", ""))
        ws.cell(row=current_row, column=2, value=f"${hotel_row.get('low', 'N/A')}")
        ws.cell(row=current_row, column=3, value=f"${hotel_row.get('high', 'N/A')}")
        ws.cell(row=current_row, column=6, value=hotel_row.get("desc", ""))
        current_row += 1

# Auto size columns
for col in ws.iter_cols():
    max_length = max((len(str(cell.value)) for cell in col if cell.value and not isinstance(cell, MergedCell)), default=10)
    if not isinstance(col[0], MergedCell):
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 50)

filename = f"golf_trip_{city.replace(' ', '_')}_{dates.replace(' ', '_')}.xlsx"
wb.save(filename)
print(f"\n✅ Spreadsheet saved as: {filename}")